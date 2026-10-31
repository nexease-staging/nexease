import bootstrap
import unittest,json,io,zipfile
from datetime import datetime
from decimal import Decimal
from unittest.mock import patch
from django.test import RequestFactory
from mongoengine.connection import get_db
from apps.organizations.models import Organization
from apps.accounts.models import User
from apps.authorization.models import Role
from apps.products.models import Product,Category
from apps.sales.models import Customer,Invoice,InvoiceItem
from apps.business.models import BusinessRecord as Record
from apps.business.billing import clean_details,tax_heads,invoice_details,invoice_settings,TEMPLATES
from apps.business.workforce import report,payslips
from apps.business.direct_entry import direct_invoice
from apps.business.print_documents import invoice_pdf
from apps.business.taxation import gst_report
from apps.business.books import ensure_accounts
from apps.sales.services.invoice_service import InvoiceService
from apps.finance.models import JournalEntry
from apps.saas import ai
from apps.saas.models import AIProposal
from apps.saas.ai_actions import confirm_extended
from pypdf import PdfReader
from openpyxl import load_workbook

class UpgradeTests(unittest.TestCase):
 def setUp(self):
  for c in get_db().list_collection_names():get_db().drop_collection(c)
  self.org=Organization(name='Example Staffing',email='test@example.com',address='Mumbai, Maharashtra',currency='INR').save()
  self.user=User(organization=self.org,email='admin@example.com',password='test',first_name='Admin',last_name='Test').save()
  self.user.role=Role(organization=self.org,name='Admin',is_system=True).save();self.user.save()
  self.cat=Category(organization=self.org,name='Services').save()
  self.product=Product(organization=self.org,category=self.cat,sku='S1',name='Labour supply',unit='DAY',item_type='service',hsn_code='998519').save()
  self.customer=Customer(organization=self.org,code='C1',name='Client').save()
  self.reg=Record(organization=self.org,created_by=self.user,kind='gst_settings',key='registration',data={'registration':'regular','gstin':'27ABCDE1234F1Z5','state_code':'27'}).save()
 def req(self,method='GET',body=None):
  r=RequestFactory().get('/',body or {}) if method=='GET' else RequestFactory().post('/',json.dumps(body or {}),content_type='application/json')
  r.api_user=self.user;r.api_organization=self.org;r.user=self.user
  return r
 def call(self,view,method='GET',body=None,allowed=True,**kwargs):
  with patch('apps.authorization.services.AuthorizationService.has_permission',return_value=allowed):return view.__wrapped__(self.req(method,body),**kwargs)
 def invoice(self,pos='27',status='DRAFT'):
  line=InvoiceItem(product=self.product,item_type='service',item_name='Labour supply',hsn_code='998519',unit='DAY',quantity=2,unit_price=100,line_subtotal=200,tax_rate=18,line_tax=36,line_total=236)
  details=clean_details({'supply_type':'taxable','place_of_supply':pos,'po_number':'PO-42'},self.org,[line])
  return Invoice(organization=self.org,customer=self.customer,created_by=self.user,invoice_number='I1',invoice_date=datetime(2026,10,1),items=[line],subtotal=200,tax_amount=36,total_amount=236,balance_due=236,billing_name='Client',billing_details=details,status=status).save()
 def test_tax_heads(self):
  inv=self.invoice();self.assertEqual(tax_heads(inv),(Decimal(18),Decimal(18),Decimal(0)))
  inv.billing_details['place_of_supply']='29';self.assertEqual(tax_heads(inv),(Decimal(0),Decimal(0),Decimal(36)))
 def test_non_taxable_rejects_tax_and_no_partial_masters(self):
  body={'invoice_number':'BAD','invoice_date':'2026-10-01','customer':{'code':'NEW','name':'New'},'items':[{'sku':'NEW','name':'New','unit':'DAY','quantity':'1','unit_price':'100','tax_rate':'18','item_type':'service','hsn_code':'998519'}],'billing_details':{'supply_type':'exempt','place_of_supply':'27'}}
  response=self.call(direct_invoice,'POST',body);self.assertEqual(response.status_code,400,response.content);self.assertFalse(Customer.objects(code='NEW'));self.assertFalse(Product.objects(sku='NEW'))
 def test_service_has_no_stock_review_and_hsn_snapshot(self):
  body={'invoice_number':'I1','invoice_date':'2026-10-01','customer':{'code':'C1','name':'Client'},'items':[{'sku':'S1','name':'Labour supply','unit':'DAY','quantity':'2','unit_price':'100','tax_rate':'18'}],'billing_details':{'supply_type':'taxable','place_of_supply':'27'}}
  response=self.call(direct_invoice,'POST',body);self.assertEqual(response.status_code,200,response.content)
  inv=Invoice.objects.get();self.assertFalse(inv.stock_review_required);self.assertEqual(inv.items[0].hsn_code,'998519')
  self.product.hsn_code='998518';self.product.save();inv.reload();self.assertEqual(inv.items[0].hsn_code,'998519')
 def test_templates_render_hsn_po_and_long_invoice(self):
  inv=self.invoice()
  for t in TEMPLATES:
   inv.billing_details['template']=t[0];pdf=invoice_pdf(inv);reader=PdfReader(io.BytesIO(pdf));text=''.join(p.extract_text() for p in reader.pages);self.assertIn('998519',text);self.assertIn('PO-42',text);self.assertIn('CGST',text)
  inv.items=inv.items*50;self.assertGreater(len(PdfReader(io.BytesIO(invoice_pdf(inv))).pages),1)
 def test_tax_export_no_draft_or_double_count(self):
  inv=self.invoice(status='ISSUED')
  response=self.call(gst_report,body={'from':'2026-10-01','to':'2026-10-31','form':'GSTR-3B'})
  wb=load_workbook(io.BytesIO(response.content));self.assertEqual(wb['Reviewed register'].max_row,2);self.assertEqual(wb['HSN SAC invoice lines'].max_row,2);self.assertEqual(wb['Reviewed register']['J2'].value,18)
  Record(organization=self.org,created_by=self.user,kind='gst_adjustments',key='outward:I1',data=dict(document_number='I1',date='2026-10-01',direction='outward',supply='taxable',party_gstin='',party_name='Client',place_of_supply='27',hsn='998519',taxable='200',cgst='18',sgst='18',igst='0',cess='0',itc_eligible=False,source_reference=str(inv.id))).save()
  response=self.call(gst_report,body={'from':'2026-10-01','to':'2026-10-31'});wb=load_workbook(io.BytesIO(response.content));self.assertEqual(wb['Reviewed register'].max_row,2)
 def test_service_and_gst_journal_balances(self):
  inv=self.invoice();ensure_accounts(self.org,self.user)
  with patch('apps.authorization.services.AuthorizationService.has_permission',return_value=True):InvoiceService._post_invoice_accounting(user=self.user,organization=self.org,invoice=inv)
  j=JournalEntry.objects(source_type='SALES_INVOICE').get();self.assertEqual(j.total_debit,j.total_credit)
  keys={l.account.system_key for l in j.lines};self.assertTrue({'SERVICE_REVENUE','OUTPUT_CGST','OUTPUT_SGST'}.issubset(keys))
 def test_details_locked_and_permission_guard(self):
  inv=self.invoice(status='ISSUED');self.assertEqual(self.call(invoice_details,'POST',{},invoice_id=str(inv.id)).status_code,400)
  self.assertEqual(self.call(invoice_details,invoice_id=str(inv.id),allowed=False).status_code,403)
  other=Organization(name='Other',email='other@example.com').save();req=self.req();req.api_organization=other
  with patch('apps.authorization.services.AuthorizationService.has_permission',return_value=True):self.assertEqual(invoice_details.__wrapped__(req,invoice_id=str(inv.id)).status_code,404)
 def test_export_includes_over_1000_and_formula_escape(self):
  Record.objects.insert([Record(organization=self.org,created_by=self.user,kind='employees',key=str(i),data={'name':'=DANGER' if i==0 else str(i),'code':str(i)}) for i in range(1002)])
  response=self.call(report,body={'format':'csv'},kind='employees');self.assertEqual(response.status_code,200);self.assertEqual(len(response.content.decode().splitlines()),1003);self.assertIn("'=DANGER",response.content.decode())
  self.assertEqual(self.call(report,kind='employees',allowed=False).status_code,403)
 def test_payslip_tenant_draft_guard_and_zip(self):
  run=Record(organization=self.org,created_by=self.user,kind='payruns',key='2026-10',data={'status':'draft','rows':[{'employee_id':'E1','employee':{'name':'Employee','code':'E1'},'result':{'gross':'100','net':'90','deductions':'10'},'inputs':{}}]}).save()
  self.assertEqual(self.call(payslips,run_id=str(run.id)).status_code,400);run.data['status']='approved';run.save()
  response=self.call(payslips,run_id=str(run.id));z=zipfile.ZipFile(io.BytesIO(response.content));self.assertEqual(len(z.namelist()),1);self.assertIn('Employee',PdfReader(io.BytesIO(z.read(z.namelist()[0]))).pages[0].extract_text())
  self.assertEqual(self.call(payslips,body={'employee_id':'OTHER'},run_id=str(run.id)).status_code,404)
 def test_ai_employee_and_category_validators(self):
  for kind in ('employee','category'):
   fields={k:'' for k in ai.FIELDS};fields.update(name='AI '+kind,joining_date='2026-10-01',basis='daily',rate='500',employment_type='site',active='true')
   proposal=AIProposal(organization_id=str(self.org.id),user_id=str(self.user.id),request_key='request-'+kind,kind=kind,state='review',data={'kind':kind,'fields':fields,'items':[]}).save()
   req=self.req('POST',{})
   with patch('apps.authorization.services.AuthorizationService.has_permission',return_value=True),patch('apps.organizations.api_context_service.APIOrganizationContextService.resolve',return_value={'user':self.user,'organization':self.org}):confirm_extended(req,proposal)
   proposal.reload();self.assertEqual(proposal.state,'saved',proposal.result)
 def test_service_rejected_by_stock_service(self):
  from apps.inventory.services.inventory_service import InventoryService
  with self.assertRaisesRegex(ValueError,'stock'):InventoryService._validate_product(self.product,self.org)

 def test_service_order_fulfils_without_inventory(self):
  from apps.inventory.models import Warehouse,Inventory,StockMovement
  from apps.sales.models import SalesOrder,SalesOrderItem
  from apps.sales.services.sales_order_service import SalesOrderService
  wh=Warehouse(organization=self.org,code='W1',name='Office').save()
  so=SalesOrder(organization=self.org,customer=self.customer,warehouse=wh,so_number='SO1',order_date=datetime(2026,10,1),created_by=self.user,items=[SalesOrderItem(product=self.product,quantity=2,unit_price=100,line_subtotal=200,line_total=200)],subtotal=200,total_amount=200).save()
  with patch('apps.authorization.services.AuthorizationService.has_permission',return_value=True):
   SalesOrderService.confirm_sales_order(user=self.user,organization=self.org,sales_order_id=str(so.id))
   SalesOrderService.fulfill_sales_order(user=self.user,organization=self.org,sales_order_id=str(so.id),raw_items=[{'product':self.product,'quantity':'2'}])
  so.reload();self.assertEqual(so.status,'FULFILLED');self.assertEqual(Inventory.objects.count(),0);self.assertEqual(StockMovement.objects.count(),0)
 def test_credit_reverses_tax_and_service_controls_with_discount(self):
  from types import SimpleNamespace
  from bson import ObjectId
  from apps.sales.services.credit_note_service import CreditNoteService
  inv=self.invoice();ensure_accounts(self.org,self.user)
  credit=SimpleNamespace(id=ObjectId(),organization=self.org,invoice=inv,credit_note_number='CN1',credit_note_date=datetime(2026,10,2),subtotal=Decimal(200),tax_amount=Decimal('34.2'),total_amount=Decimal('224.2'),items=[SimpleNamespace(product=self.product,line_subtotal=Decimal(200),line_tax=Decimal('34.2'),line_total=Decimal('224.2'))])
  with patch('apps.authorization.services.AuthorizationService.has_permission',return_value=True):CreditNoteService._post_credit_note_accounting(user=self.user,organization=self.org,credit_note=credit)
  j=JournalEntry.objects(source_type='SALES_CREDIT_NOTE').get();self.assertEqual(j.total_debit,Decimal('224.2'));self.assertEqual(j.total_credit,Decimal('224.2'));self.assertIn('SERVICE_REVENUE',{l.account.system_key for l in j.lines})
