import bootstrap
import unittest,json
from datetime import datetime
from decimal import Decimal
from unittest.mock import patch
from django.test import RequestFactory
from mongoengine.connection import get_db
from apps.organizations.models import Organization
from apps.accounts.models import User
from apps.business.models import BusinessRecord as Record
from apps.business.views import records, payroll, validate
from apps.business.payroll import calculate
from apps.business.direct_entry import direct_invoice
from apps.business.purchase_entry import quick_purchase
from apps.sales.models import Invoice
from apps.purchasing.models import PurchaseOrder,VendorBill
from apps.migration.services.migration_schema_service import MigrationSchemaService
from apps.migration.services.financial_import import validate_financial

class ReleaseTests(unittest.TestCase):
 def setUp(self):
  for c in get_db().list_collection_names():get_db().drop_collection(c)
  self.org=Organization(name='Example',email='example@example.com').save();self.other=Organization(name='Other',email='other@example.com').save()
  self.user=User(organization=self.org,email='user@example.com',password='not-a-real-password',first_name='Test',last_name='User').save()
  self.factory=RequestFactory()
 def req(self,method='get',data=None):
  r=getattr(self.factory,method)('/',data=json.dumps(data or {}),content_type='application/json') if method!='get' else self.factory.get('/')
  r.api_organization=self.org;r.api_user=self.user;return r
 def call(self,view,r,*args):
  # The existing authentication decorator is separately covered by the original suite.
  with patch('apps.authorization.services.AuthorizationService.has_permission',return_value=True):return view.__wrapped__(r,*args)
 def employee(self,code='E1'):
  d={'code':code,'name':'Worker','joining_date':'2026-01-01','basis':'monthly','rate':'30000','active':True}
  return Record(organization=self.org,kind='employees',key=code,data=validate('employees',d,self.org),created_by=self.user).save()
 def test_monthly_pay(self):
  r=calculate({'basis':'monthly','rate':'30000'},{'period_days':30,'paid_days':15},{})
  self.assertEqual(r['net'],'15000.00')
 def test_daily_hourly_piece(self):
  for basis,inputs,expected in [('daily',{'paid_days':5},'500.00'),('hourly',{'hours':8},'800.00'),('piece',{'units':12},'1200.00')]:self.assertEqual(calculate({'basis':basis,'rate':100},inputs,{})['net'],expected)
 def test_statutory_policy_required(self):
  with self.assertRaises(ValueError):calculate({'basis':'monthly','rate':30000,'pf_member':True},{'period_days':30,'paid_days':30,'pf_wages':15000},{})
 def test_pf_esi_rounding(self):
  r=calculate({'basis':'monthly','rate':10001,'pf_member':True,'esi_member':True},{'period_days':30,'paid_days':30,'pf_wages':10001,'esi_wages':10001},{'confirmed':True,'pf_cap':15000,'pf_employee_rate':12,'pf_employer_rate':12,'esi_employee_rate':.75,'esi_employer_rate':3.25})
  self.assertEqual(r['pf'],'1200.12');self.assertEqual(r['esi'],'76.00');self.assertEqual(r['employer_esi'],'326.00')
 def test_invalid_days(self):
  with self.assertRaises(ValueError):calculate({'basis':'monthly','rate':100},{'period_days':30,'paid_days':31},{})
 def test_excess_deductions(self):
  with self.assertRaises(ValueError):calculate({'basis':'daily','rate':100},{'paid_days':1,'tds':101},{})
 def test_nonfinite(self):
  with self.assertRaises(ValueError):calculate({'basis':'daily','rate':'NaN'},{},{})
 def test_records_tenant_isolation(self):
  Record(organization=self.other,kind='sites',key='secret',data={'name':'Secret'},created_by=self.user).save()
  response=self.call(records,self.req(),'sites');self.assertEqual(json.loads(response.content)['data']['records'],[])
 def test_permission_denied(self):
  with patch('apps.authorization.services.AuthorizationService.has_permission',return_value=False):self.assertEqual(records.__wrapped__(self.req(),'employees').status_code,403)
 def test_stale_edit(self):
  emp=self.employee();r=self.call(records,self.req('patch',{'data':emp.data,'revision':999}),'employees',str(emp.id));self.assertEqual(r.status_code,400)
 def test_cross_tenant_assignment(self):
  emp=Record(organization=self.other,kind='employees',key='E',data={},created_by=self.user).save()
  with self.assertRaises(ValueError):validate('assignments',{'employee_id':str(emp.id),'site_id':'000000000000000000000000','start_date':'2026-01-01','end_date':'2026-02-01','shift':'Day','status':'active'},self.org)
 def test_assignment_overlap(self):
  emp=self.employee();site=Record(organization=self.org,kind='sites',key='S',data={},created_by=self.user).save();d={'employee_id':str(emp.id),'site_id':str(site.id),'start_date':'2026-01-01','end_date':'2026-02-01','shift':'Day','status':'active'}
  Record(organization=self.org,kind='assignments',key='A',data=d,created_by=self.user).save()
  with self.assertRaises(ValueError):validate('assignments',d,self.org)
 def test_payroll_lifecycle(self):
  emp=self.employee();d=validate('attendance',{'employee_id':str(emp.id),'month':'2026-10','period_days':'31','paid_days':'31'},self.org)
  Record(organization=self.org,kind='attendance',key=str(emp.id)+':2026-10',data=d,created_by=self.user).save()
  r=self.call(payroll,self.req('post',{'month':'2026-10'}));self.assertEqual(r.status_code,200,r.content)
  run=json.loads(r.content)['data']['record'];r=self.call(payroll,self.req('post',{'month':'2026-10','action':'approve','revision':run['revision']}));self.assertEqual(r.status_code,200,r.content)
  with self.assertRaises(ValueError):validate('attendance',d,self.org)
  self.assertEqual(self.call(payroll,self.req('post',{'month':'2026-10'})).status_code,400)
 def direct_payload(self,number='D1'):
  return {'invoice_number':number,'invoice_date':'2026-10-21','customer':{'code':'C1','name':'Customer'},'items':[{'sku':'SKU1','name':'Widget','unit':'PCS','quantity':'2','unit_price':'100','tax_rate':'18'}]}
 def test_multiple_direct_invoices_and_masters(self):
  for number in ['D1','D2']:
   r=self.call(direct_invoice,self.req('post',self.direct_payload(number)));self.assertEqual(r.status_code,200,r.content)
  self.assertEqual(Invoice.objects.count(),2);self.assertEqual(Invoice.objects.first().total_amount,Decimal('236.00'))
  self.assertIsNone(Invoice.objects.first().sales_order)
 def test_invalid_direct_no_masters(self):
  p=self.direct_payload();p['items'][0]['quantity']='-1'
  self.assertEqual(self.call(direct_invoice,self.req('post',p)).status_code,400)
  from apps.sales.models import Customer
  self.assertEqual(Customer.objects.count(),0)
 def test_purchase_entry(self):
  for kind in ['bill','purchase_order']:
   p=self.direct_payload();p.update(kind=kind,number=kind,date='2026-10-21',supplier={'code':'S1','name':'Supplier'})
   r=self.call(quick_purchase,self.req('post',p));self.assertEqual(r.status_code,200,r.content)
  self.assertEqual(PurchaseOrder.objects.count(),1);self.assertEqual(VendorBill.objects.count(),1)
 def test_new_schemas(self):
  for kind in ['VENDOR_BILLS','CUSTOMER_PAYMENTS','SUPPLIER_PAYMENTS']:self.assertTrue(MigrationSchemaService.get_schema(kind)['fields'])
 def test_payment_validation_rejects_missing_invoice(self):
  r=validate_financial(self.org,'CUSTOMER_PAYMENTS',{'invoice_number':'missing','amount':'5','payment_method':'CASH','bank_account_id':'000000000000000000000000','payment_date':'2026-10-21','reference_number':'R1'},set());self.assertFalse(r['valid'])
 def test_template_download_has_usable_sheets(self):
  from apps.migration.api.v1.views import migration_template_api
  from openpyxl import load_workbook
  from io import BytesIO
  r=self.call(migration_template_api,self.req(),'SUPPLIER_PAYMENTS')
  self.assertEqual(r.status_code,200)
  wb=load_workbook(BytesIO(r.content));self.assertIn('Example only',wb.sheetnames)
  self.assertIn('bank_account_id',[x.value for x in wb['Data'][1]])
 def test_gst_workpaper_marks_unclassified_sources(self):
  from apps.business.taxation import gst_report
  self.call(direct_invoice,self.req('post',self.direct_payload()))
  Invoice.objects.update(set__status='ISSUED')
  r=self.req();r.GET=r.GET.copy();r.GET.update({'from':'2026-10-01','to':'2026-10-31','form':'GSTR-1'})
  response=self.call(gst_report,r)
  from openpyxl import load_workbook
  from io import BytesIO
  wb=load_workbook(BytesIO(response.content));self.assertEqual(wb['Source reconciliation']['I2'].value,'NEEDS CLASSIFICATION')
  self.assertEqual(wb['Reviewed register'].max_row,1)
 def test_signup_scope_saved(self):
  from django.core.management import call_command
  from apps.accounts.organization_signup_service import OrganizationSignupService
  call_command('seed_permissions',verbosity=0,stdout=__import__('io').StringIO())
  OrganizationSignupService.create_organization_account(organization_name='Staffing',organization_email='staffing@example.com',first_name='A',last_name='B',user_email='new@example.com',password='ExampleOnly#82933',business_scope='labour_supply')
  self.assertEqual(Organization.objects(name='Staffing').first().business_scope,'labour_supply')
 def test_duplicate_lines_do_not_create_masters(self):
  p=self.direct_payload();p['items']*=2
  self.assertEqual(self.call(direct_invoice,self.req('post',p)).status_code,400)
  from apps.sales.models import Customer
  self.assertEqual(Customer.objects.count(),0)
 def test_worker_lookup_excludes_pay_and_bank(self):
  from apps.business.views import worker_lookup
  self.employee();response=self.call(worker_lookup,self.req())
  d=json.loads(response.content)['data']['records'][0]['data']
  self.assertEqual(set(d),{'name','code','active','employment_type'})

if __name__=='__main__':unittest.main(verbosity=2)
