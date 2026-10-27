from csv import DictReader
from io import BytesIO, StringIO
from pathlib import Path

from openpyxl import load_workbook


class MigrationFileParserError(
    ValueError,
):
    pass


class MigrationFileParser:

    SUPPORTED_EXTENSIONS = {
        ".csv",
        ".xlsx",
    }

    MAX_SAMPLE_ROWS = 20

    @staticmethod
    def normalize_header(
        value,
    ):
        if value is None:
            return ""

        return " ".join(
            str(value)
            .strip()
            .split()
        )

    @classmethod
    def parse(
        cls,
        *,
        filename,
        content,
        sample_limit=10,
    ):
        filename = str(
            filename or ""
        ).strip()

        if not filename:
            raise MigrationFileParserError(
                "Filename is required."
            )

        if not isinstance(
            content,
            bytes,
        ):
            raise MigrationFileParserError(
                "File content must be bytes."
            )

        if not content:
            raise MigrationFileParserError(
                "Migration file is empty."
            )

        extension = (
            Path(filename)
            .suffix
            .lower()
        )

        if (
            extension
            not in
            cls.SUPPORTED_EXTENSIONS
        ):
            raise MigrationFileParserError(
                "Unsupported migration file type."
            )

        if (
            not isinstance(sample_limit, int)
            or isinstance(sample_limit, bool)
            or sample_limit < 1
            or sample_limit > cls.MAX_SAMPLE_ROWS
        ):
            raise MigrationFileParserError(
                "Sample limit must be between "
                f"1 and {cls.MAX_SAMPLE_ROWS}."
            )

        if extension == ".csv":
            result = cls._parse_csv(
                content=content,
                sample_limit=sample_limit,
            )
        else:
            result = cls._parse_xlsx(
                content=content,
                sample_limit=sample_limit,
            )

        return {
            "filename": filename,
            "extension": extension,
            **result,
        }

    @classmethod
    def read_rows(
        cls,
        *,
        filename,
        content,
    ):
        filename = str(
            filename or ""
        ).strip()

        if not filename:
            raise MigrationFileParserError(
                "Filename is required."
            )

        if not isinstance(
            content,
            bytes,
        ):
            raise MigrationFileParserError(
                "File content must be bytes."
            )

        if not content:
            raise MigrationFileParserError(
                "Migration file is empty."
            )

        extension = (
            Path(filename)
            .suffix
            .lower()
        )

        if (
            extension
            not in
            cls.SUPPORTED_EXTENSIONS
        ):
            raise MigrationFileParserError(
                "Unsupported migration file type."
            )

        if extension == ".csv":
            return cls._read_csv_rows(
                content=content,
            )

        return cls._read_xlsx_rows(
            content=content,
        )

    @classmethod
    def _parse_csv(
        cls,
        *,
        content,
        sample_limit,
    ):
        try:
            text = content.decode(
                "utf-8-sig"
            )
        except UnicodeDecodeError as exc:
            raise MigrationFileParserError(
                "CSV file must use UTF-8 encoding."
            ) from exc

        reader = DictReader(
            StringIO(text)
        )

        if reader.fieldnames is None:
            raise MigrationFileParserError(
                "CSV file does not contain "
                "a header row."
            )

        raw_headers = list(
            reader.fieldnames
        )

        headers = (
            cls._validate_headers(
                raw_headers
            )
        )

        sample_rows = []
        total_rows = 0

        for row in reader:
            if cls._row_is_blank(
                row.values()
            ):
                continue

            total_rows += 1

            if (
                len(sample_rows)
                < sample_limit
            ):
                sample_rows.append(
                    cls._serialize_row(
                        headers=headers,
                        values=[
                            row.get(header)
                            for header
                            in raw_headers
                        ],
                    )
                )

        return {
            "sheet_name": None,
            "columns": headers,
            "total_rows": total_rows,
            "sample_rows": sample_rows,
        }

    @classmethod
    def _parse_xlsx(
        cls,
        *,
        content,
        sample_limit,
    ):
        try:
            workbook = load_workbook(
                filename=BytesIO(content),
                read_only=True,
                data_only=True,
            )
        except Exception as exc:
            raise MigrationFileParserError(
                "Unable to read XLSX file."
            ) from exc

        try:
            worksheet = workbook.active

            rows = worksheet.iter_rows(
                values_only=True
            )

            try:
                raw_headers = list(
                    next(rows)
                )
            except StopIteration as exc:
                raise MigrationFileParserError(
                    "XLSX file is empty."
                ) from exc

            headers = (
                cls._validate_headers(
                    raw_headers
                )
            )

            sample_rows = []
            total_rows = 0

            for values in rows:
                if cls._row_is_blank(
                    values
                ):
                    continue

                total_rows += 1

                if (
                    len(sample_rows)
                    < sample_limit
                ):
                    sample_rows.append(
                        cls._serialize_row(
                            headers=headers,
                            values=values,
                        )
                    )

            return {
                "sheet_name":
                    worksheet.title,
                "columns": headers,
                "total_rows": total_rows,
                "sample_rows": sample_rows,
            }

        finally:
            workbook.close()

    @classmethod
    def _read_csv_rows(
        cls,
        *,
        content,
    ):
        try:
            text = content.decode(
                "utf-8-sig"
            )
        except UnicodeDecodeError as exc:
            raise MigrationFileParserError(
                "CSV file must use UTF-8 encoding."
            ) from exc

        reader = DictReader(
            StringIO(text)
        )

        if reader.fieldnames is None:
            raise MigrationFileParserError(
                "CSV file does not contain "
                "a header row."
            )

        raw_headers = list(
            reader.fieldnames
        )

        headers = cls._validate_headers(
            raw_headers
        )

        rows = []

        for row in reader:
            values = [
                row.get(header)
                for header
                in raw_headers
            ]

            if cls._row_is_blank(
                values
            ):
                continue

            rows.append(
                cls._serialize_row(
                    headers=headers,
                    values=values,
                )
            )

        return {
            "columns": headers,
            "rows": rows,
            "total_rows": len(rows),
            "sheet_name": None,
        }

    @classmethod
    def _read_xlsx_rows(
        cls,
        *,
        content,
    ):
        try:
            workbook = load_workbook(
                filename=BytesIO(content),
                read_only=True,
                data_only=True,
            )
        except Exception as exc:
            raise MigrationFileParserError(
                "Unable to read XLSX file."
            ) from exc

        try:
            worksheet = workbook.active

            iterator = worksheet.iter_rows(
                values_only=True
            )

            try:
                raw_headers = list(
                    next(iterator)
                )
            except StopIteration as exc:
                raise MigrationFileParserError(
                    "XLSX file is empty."
                ) from exc

            headers = cls._validate_headers(
                raw_headers
            )

            parsed_rows = []

            for values in iterator:
                if cls._row_is_blank(
                    values
                ):
                    continue

                parsed_rows.append(
                    cls._serialize_row(
                        headers=headers,
                        values=values,
                    )
                )

            return {
                "columns": headers,
                "rows": parsed_rows,
                "total_rows":
                    len(parsed_rows),
                "sheet_name":
                    worksheet.title,
            }

        finally:
            workbook.close()

    @classmethod
    def _validate_headers(
        cls,
        raw_headers,
    ):
        if not raw_headers:
            raise MigrationFileParserError(
                "Migration file does not "
                "contain columns."
            )

        headers = [
            cls.normalize_header(
                value
            )
            for value in raw_headers
        ]

        blank_positions = [
            index + 1
            for index, header
            in enumerate(headers)
            if not header
        ]

        if blank_positions:
            raise MigrationFileParserError(
                "Blank column headers found at "
                "positions: "
                +
                ", ".join(
                    str(position)
                    for position
                    in blank_positions
                )
                +
                "."
            )

        normalized_lookup = {}

        for header in headers:
            key = header.casefold()

            normalized_lookup.setdefault(
                key,
                [],
            ).append(header)

        duplicates = [
            values[0]
            for values
            in normalized_lookup.values()
            if len(values) > 1
        ]

        if duplicates:
            raise MigrationFileParserError(
                "Duplicate column headers found: "
                +
                ", ".join(
                    duplicates
                )
                +
                "."
            )

        return headers

    @staticmethod
    def _row_is_blank(
        values,
    ):
        return all(
            value is None
            or (
                isinstance(value, str)
                and not value.strip()
            )
            for value in values
        )

    @classmethod
    def _serialize_row(
        cls,
        *,
        headers,
        values,
    ):
        values = list(values)

        result = {}

        for index, header in enumerate(
            headers
        ):
            value = (
                values[index]
                if index < len(values)
                else None
            )

            result[header] = (
                cls._serialize_value(
                    value
                )
            )

        return result

    @staticmethod
    def _serialize_value(
        value,
    ):
        if value is None:
            return None

        if hasattr(
            value,
            "isoformat",
        ):
            try:
                return value.isoformat()
            except (
                TypeError,
                ValueError,
            ):
                pass

        if isinstance(
            value,
            (
                str,
                int,
                float,
                bool,
            ),
        ):
            return value

        return str(value)