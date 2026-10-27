import calendar, json, re
from datetime import datetime, date
from decimal import Decimal
from io import BytesIO
from django.http import HttpResponse
from django.views.decorators.http import require_http_methods
from mongoengine.errors import ValidationError, NotUniqueError
from openpyxl import Workbook
from apps.core.api.decorators import api_login_required
from apps.authorization.services import AuthorizationService
from apps.core.services.api_response_service import APIResponseService as R
from apps.business.models import BusinessRecord as Record, BusinessAudit
from apps.business.schema import SCHEMAS
from apps.business.profiles import PROFILES
from apps.business.payroll import amount, calculate, money
from .automation import guard_source,write_lock,run_data

def permitted(request, code):return AuthorizationService.has_permission(request.api_user,code)
def serialize(x):return {'id':str(x.id),'key':x.key,'revision':x.revision,'data':x.data,'created_at':x.created_at.isoformat()}
def payload(request):
    try:d=json.loads(request.body)
    except (ValueError,UnicodeDecodeError):raise ValueError('Invalid JSON body.')
    if not isinstance(d,dict):raise ValueError('Expected an object.')
    return d

def validate(kind, data, org, record_id=None):
    schema=SCHEMAS[kind]['fields'];clean={}
    if not isinstance(data,dict) or set(data)-set(schema):raise ValueError('Unknown fields supplied.')
    for key, spec in schema.items():
        value=data.get(key)
        if value in (None,''):
            if spec['required']:raise ValueError(spec['label']+' is required.')
            clean[key]=False if spec['type']=='boolean' else ('0' if spec['type']=='number' else '')
            continue
        t=spec['type']
        if t=='boolean':
            if not isinstance(value,bool):raise ValueError(spec['label']+' must be true or false.')
        elif t=='number':value=str(amount(value,spec['label']))
        else:
            if not isinstance(value,str) or len(value)>1000:raise ValueError(spec['label']+' is invalid.')
            value=value.strip()
            if t=='date':date.fromisoformat(value)
            if t=='month' and not re.fullmatch(r'\d{4}-(0[1-9]|1[0-2])',value):raise ValueError('Use YYYY-MM.')
            if t=='select' and value not in spec['choices']:raise ValueError('Invalid '+spec['label'])
            if t in ('employees','sites') and not Record.objects(id=value,organization=org,kind=t).first():raise ValueError('Related record not found in this organization.')
        clean[key]=value
    for a,b in [('start_date','end_date'),('joining_date','end_date'),('effective_from','effective_to')]:
        if clean.get(a) and clean.get(b) and clean[b]<clean[a]:raise ValueError('End date precedes start date.')
    guard_source(kind,clean,org,record_id)
    if kind=='timesheets':
        if not clean['site_id']:
            employee=Record.objects(id=clean['employee_id'],organization=org,kind='employees').first()
            if not employee or employee.data.get('employment_type')!='inhouse':raise ValueError('Choose a client site for a site worker.')
            clean['billable_units']='0';clean['billing_rate']='0'
        if amount(clean['hours'])+amount(clean['overtime_hours'])>24:raise ValueError('Daily hours exceed 24.')
        assignments=Record.objects(organization=org,kind='assignments',data__employee_id=clean['employee_id'],data__site_id=clean['site_id'])
        if clean['site_id'] and not any(x.data['status']!='cancelled' and x.data['start_date']<=clean['date']<=x.data['end_date'] and x.data['shift'].casefold()==clean['shift'].casefold() for x in assignments):raise ValueError('Create a matching worker/site/shift assignment for this date first.')
        other_hours=sum(amount(x.data['hours'])+amount(x.data['overtime_hours']) for x in Record.objects(organization=org,kind='timesheets',data__employee_id=clean['employee_id'],data__date=clean['date']) if str(x.id)!=str(record_id))
        if other_hours+amount(clean['hours'])+amount(clean['overtime_hours'])>24:raise ValueError('Combined site hours exceed 24 for this worker/date.')
    if kind in ('advances','expenses') and amount(clean['amount'])<=0:raise ValueError('Amount must be positive.')
    if kind=='advances' and amount(clean['recovered'])>amount(clean['amount']):raise ValueError('Recovery exceeds the advance.')
    if kind=='leave_requests' and amount(clean['days'])>(date.fromisoformat(clean['end_date'])-date.fromisoformat(clean['start_date'])).days+1:raise ValueError('Leave days exceed the date range.')
    if kind=='assignments' and clean['status']=='active':
        for x in Record.objects(organization=org,kind=kind):
            d=x.data
            if str(x.id)!=str(record_id) and d.get('status')=='active' and d.get('employee_id')==clean['employee_id'] and d.get('shift','').casefold()==clean['shift'].casefold() and d['start_date']<=clean['end_date'] and d['end_date']>=clean['start_date']:raise ValueError('This worker already has an overlapping assignment in this shift.')
    if kind=='attendance':
        y,m=map(int,clean['month'].split('-'));days=calendar.monthrange(y,m)[1]
        if not 1<=amount(clean['period_days'])<=days or amount(clean['paid_days'])>amount(clean['period_days']):raise ValueError('Review paid days and period days.')
        if Record.objects(organization=org,kind='payruns',key=clean['month'],data__status__in=['approved','paid']).first():raise ValueError('This payroll month is locked.')
    if kind=='policy':
        if any(amount(clean[k])>100 for k in clean if k.endswith('_rate')):raise ValueError('Rates cannot exceed 100%.')
        for x in Record.objects(organization=org,kind=kind):
            if str(x.id)!=str(record_id) and x.data['effective_from']<=clean['effective_to'] and x.data['effective_to']>=clean['effective_from']:raise ValueError('Payroll policies cannot overlap.')
    if kind=='gst_settings':
        if not re.fullmatch(r'\d{2}',clean['state_code']):raise ValueError('State code must have two digits.')
        if clean['registration']!='unregistered' and not re.fullmatch(r'\d{2}[A-Z]{5}\d{4}[A-Z][A-Z0-9]Z[A-Z0-9]',clean['gstin']):raise ValueError('Check GSTIN format.')
    if kind=='gst_adjustments' and amount(clean['igst']) and (amount(clean['cgst']) or amount(clean['sgst'])):raise ValueError('Use IGST or CGST/SGST, not both.')
    return clean

def key_for(kind,d):
    if kind in ('employees','sites'):return d['code'].upper()
    if kind=='attendance':return d['employee_id']+':'+d['month']
    if kind=='timesheets':return d['employee_id']+':'+d['date']+':'+d['shift'].casefold()
    if kind in ('expenses','advances'):
        from hashlib import sha256
        return sha256((d.get('employee_id',d.get('payee',''))+'|'+d['date']+'|'+d['reference']).encode()).hexdigest()
    if kind=='policy':return d['effective_from']
    if kind=='gst_settings':return 'registration'
    if kind=='gst_adjustments':return d['direction']+':'+d['document_number']
    import uuid
    return uuid.uuid4().hex

def audit(req,kind,id,action):BusinessAudit(organization=req.api_organization,actor=req.api_user,kind=kind,record_id=str(id),action=action).save()

@api_login_required
@require_http_methods(['GET'])
def configuration(request):
    scope=getattr(request.api_organization,'business_scope','mixed')
    return R.success(data={'schemas':{k:v for k,v in SCHEMAS.items() if permitted(request,v['permission']+'.read')},'profiles':PROFILES,'scope':scope,'profile':PROFILES.get(scope,PROFILES['mixed']),'modules':{'workforce':request.api_organization.workforce_enabled,'expenses':request.api_organization.expenses_enabled}},request=request)

@api_login_required
@require_http_methods(['GET','POST','PATCH'])
@write_lock
def records(request,kind,record_id=None):
    if kind not in SCHEMAS:return R.not_found(message='Unknown collection.',request=request)
    perm=SCHEMAS[kind]['permission']+('.read' if request.method=='GET' else '.manage')
    if not permitted(request,perm):return R.forbidden(message='Permission denied.',request=request)
    org=request.api_organization
    try:
        if request.method=='GET':
            query=Record.objects(organization=org,kind=kind).order_by('-created_at')
            term=request.GET.get('search','').strip()[:100]
            if term:
                from mongoengine.queryset.visitor import Q
                match=Q(key__icontains=term)
                for field,spec in SCHEMAS[kind]['fields'].items():
                    if spec['type'] in ('text','date','select','month'):match |= Q(**{'data__'+field+'__icontains':term})
                    if spec['type'] in ('employees','sites'):
                        ids=[str(x.id) for x in Record.objects(organization=org,kind=spec['type']).filter(Q(data__name__icontains=term)|Q(data__code__icontains=term))]
                        match |= Q(**{'data__'+field+'__in':ids})
                query=query.filter(match)
            return R.success(data={'records':[serialize(x) for x in query.limit(1000)],'total':query.count()},request=request)
        p=payload(request);raw=p.get('data')
        if kind in ('advances','expenses'):
            from .books import ensure_accounts
            ensure_accounts(org,request.api_user)
            if record_id:
                from apps.finance.models import JournalEntry
                if JournalEntry.objects(organization=org,source_id=record_id,status__in=['POSTED','REVERSED']).first():raise ValueError('This record is linked to accounting. Use an accountant-reviewed correction journal rather than editing the source.')
        if kind=='sites' and isinstance(raw,dict) and raw.get('customer_code'):
            from apps.sales.models import Customer
            c=Customer.objects(organization=org,code=raw['customer_code'],is_active=True).first()
            if not c:raise ValueError('Choose an active client.')
            raw={**raw,'address':raw.get('address') or c.billing_address,'phone':raw.get('phone') or c.phone}
        if kind=='assignments' and isinstance(raw,dict) and raw.get('billing_rate') in (None,''):
            site=Record.objects(id=raw.get('site_id'),organization=org,kind='sites').first()
            if site:raw={**raw,'billing_rate':site.data.get('billing_rate','0')}
        if request.method=='POST' and kind in ('employees','sites') and isinstance(raw,dict):
            from .numbering import choice,allocate,settings
            nk='Employee' if kind=='employees' else 'Site';selection=choice.get()
            if selection and selection.get('kind')==nk:
                code=allocate(org,nk,selection.get('number','') if selection.get('mode')=='manual' else None)
                for _ in range(100):
                    if not Record.objects(organization=org,kind=kind,key=code.upper()).first():break
                    if selection.get('mode')=='manual':raise ValueError('Code already exists.')
                    code=allocate(org,nk)
                else:raise ValueError('Choose a higher next number.')
                raw={**raw,'code':code}
            elif not raw.get('code'):
                if settings(org,nk)['mode']=='manual':raise ValueError('Enter a code.')
                raw={**raw,'code':allocate(org,nk)}
        d=validate(kind,raw,org,record_id)
        if kind in ('advances','expenses'):
            from .books import payment_account
            payment_account(org,d.get('payment_mode') or 'Bank transfer',d.get('payment_account',''))
        if request.method=='POST':
            x=Record(organization=org,kind=kind,key=key_for(kind,d),data=d,created_by=request.api_user).save()
        else:
            x=Record.objects(id=record_id,organization=org,kind=kind).first()
            if not x:return R.not_found(message='Record not found.',request=request)
            from apps.saas.entitlements import check_capacity,org_lock
            with org_lock(org.id):
                if kind=='employees' and d.get('active') and not x.data.get('active'):check_capacity(org,'employees')
                newkey=key_for(kind,d) if kind!='assignments' else x.key
                x=Record.objects(id=x.id,organization=org,revision=p.get('revision')).modify(new=True,set__data=d,set__key=newkey,inc__revision=1,set__updated_at=datetime.utcnow())
            if not x:return R.validation_error(message='Record changed. Reload before editing.',request=request)
        audit(request,kind,x.id,request.method)
        if kind in ('advances','expenses'):
            from .books import sync_record
            try:sync_record(x,request.api_user)
            except Exception as exc:
                return R.validation_error(message='Record saved but books sync needs attention. Do not enter it again. Use Settings > Business modules > Retry pending books sync. '+str(exc),request=request)
        return R.success(data={'record':serialize(x)},request=request)
    except (ValueError,ValidationError,NotUniqueError) as e:return R.validation_error(message='Duplicate record.' if isinstance(e,NotUniqueError) else str(e),request=request)

@api_login_required
@require_http_methods(['GET','POST'])
def payroll(request):
    from .payroll_api import payroll as handler
    return handler.__wrapped__(request)

def workbook_response(sheets,filename):
    wb=Workbook();wb.remove(wb.active)
    from openpyxl.styles import Font, PatternFill
    for title,rows in sheets.items():
        ws=wb.create_sheet(title[:31]);ws.freeze_panes='A2'
        for row in rows:
            ws.append([("'"+v if isinstance(v,str) and v.startswith(('=','+','-','@')) else v) for v in row])
        for cell in ws[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='346B40')
        for col in ws.columns:ws.column_dimensions[col[0].column_letter].width=min(45,max(16,max(len(str(c.value or '')) for c in col)+2))
    stream=BytesIO();wb.save(stream)
    response=HttpResponse(stream.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet');response['Content-Disposition']=f'attachment; filename="{filename}.xlsx"';return response

@api_login_required
@require_http_methods(['GET'])
def payroll_export(request,run_id):
    if not permitted(request,'payroll.read'):return R.forbidden(message='Permission denied.',request=request)
    try:run=Record.objects(id=run_id,organization=request.api_organization,kind='payruns').first()
    except ValidationError:run=None
    if not run:return R.not_found(message='Payroll not found.',request=request)
    rows=[['Code','Employee','Gross','PF employee','ESI employee','Deductions','Net','Employer PF','Employer ESI','Status','Paid','Salary outstanding','Advance outstanding']]
    bank=[['Employee','Bank account','IFSC','Net payable','Payroll month','Status']]
    for x in run_data(run)['data']['rows']:
        e=x['employee'];r=x['result'];rows.append([e['code'],e['name'],*[float(r[k]) for k in ['gross','pf','esi','deductions','net','employer_pf','employer_esi']],run.data['status'],x['paid'],x['outstanding'],x['advance_outstanding']]);bank.append([e['name'],e.get('bank_account',''),e.get('ifsc',''),float(r['net']),run.key,run.data['status']])
    return workbook_response({'Payroll register':rows,'Payment workpaper':bank,'Read me':[['Payroll export'],['Bank-specific upload format and statutory returns require separate review.'],['Marking paid records an external payment; it does not transfer funds or post a bank journal.']]},'NexEase-Payroll-'+run.key)

@api_login_required
@require_http_methods(['GET'])
def labour_export(request):
    if not permitted(request,'labour.read'):return R.forbidden(message='Permission denied.',request=request)
    try:
        start=date.fromisoformat(request.GET.get('from',''));end=date.fromisoformat(request.GET.get('to',''))
        if end<start:raise ValueError('Invalid date range.')
        rows=[['Date','Worker','Site','Shift','Hours','Overtime hours','Billable units','Rate','Billable value']];totals={}
        for x in Record.objects(organization=request.api_organization,kind='timesheets',data__status='approved'):
            d=x.data
            if not start.isoformat()<=d['date']<=end.isoformat():continue
            employee=Record.objects(organization=request.api_organization,id=d['employee_id'],kind='employees').first();site=Record.objects(organization=request.api_organization,id=d['site_id'],kind='sites').first() if d.get('site_id') else None
            value=amount(d['billable_units'])*amount(d['billing_rate']);key=d['site_id'];totals[key]=totals.get(key,Decimal(0))+value
            rows.append([d['date'],employee.data['name'] if employee else d['employee_id'],site.data['name'] if site else key,d['shift'],float(amount(d['hours'])),float(amount(d['overtime_hours'])),float(amount(d['billable_units'])),float(amount(d['billing_rate'])),float(value)])
        summary=[['Site ID','Billable value before GST']]+[[k,float(v)] for k,v in totals.items()]
        return workbook_response({'Approved timesheets':rows,'Site billing summary':summary,'Read me':[['Labour billing workpaper'],['Only approved timesheets are included. This does not create a sales invoice or payroll attendance.'],['Reconcile approved timesheets, leave and advance records to monthly payroll inputs before approval.']]},'NexEase-Labour-Billing')
    except ValueError as e:return R.validation_error(message=str(e),request=request)

@api_login_required
@require_http_methods(['GET'])
def worker_lookup(request):
    if not (permitted(request,'labour.read') or permitted(request,'payroll.read')):return R.forbidden(message='Permission denied.',request=request)
    return R.success(data={'records':[{'id':str(x.id),'data':{'name':x.data.get('name'),'code':x.data.get('code'),'employment_type':x.data.get('employment_type','site'),'active':x.data.get('active',True)}} for x in Record.objects(organization=request.api_organization,kind='employees').limit(1000)]},request=request)
