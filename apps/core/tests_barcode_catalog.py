import json
from io import BytesIO
from unittest.mock import patch
from uuid import uuid4
from django.test import SimpleTestCase, RequestFactory
from openpyxl import load_workbook
from apps.products.models import Product, Category
from apps.organizations.models import Organization
from apps.accounts.models import User
from apps.business.barcode_catalog import catalog, export_catalog, generate_missing
from apps.business.models import BusinessAudit

class BarcodeCatalogTests(SimpleTestCase):
    def setUp(self):
        self.org=Organization(name='Barcode test',email=f'{uuid4()}@example.test').save()
        self.other=Organization(name='Other',email=f'{uuid4()}@example.test').save()
        self.category=Category(organization=self.org,name='Test').save()
        self.user=User(organization=self.org,email=f'{uuid4()}@example.test',password='test',first_name='Test',last_name='User').save()
    def tearDown(self):
        Product.objects(organization__in=[self.org,self.other]).delete()
        BusinessAudit.objects(organization=self.org).delete()
        self.category.delete();self.user.delete();self.org.delete();self.other.delete()
    def item(self,**kwargs):
        return Product(organization=kwargs.pop('organization',self.org),sku=kwargs.pop('sku',uuid4().hex),name=kwargs.pop('name','Example'),category=self.category,unit='piece',**kwargs).save()
    def request(self,method='get',path='/api/v1/business/barcode/catalog/'):
        req=getattr(RequestFactory(),method)(path)
        req.api_user=self.user;req.api_organization=self.org
        return req
    def test_auto_barcode_stable_and_supplied_leading_zeroes_preserved(self):
        item=self.item();code=item.barcode
        self.assertEqual(code,Product.automatic_barcode(item.id))
        item.name='Changed';item.save();self.assertEqual(item.barcode,code)
        supplied=self.item(barcode='000123456789');self.assertEqual(supplied.barcode,'000123456789')
        self.assertNotEqual(item.barcode,self.item().barcode)
    @patch('apps.business.barcode_catalog.permitted',return_value=True)
    def test_list_is_tenant_scoped_and_includes_inactive_items(self,_):
        item=self.item(is_active=False);self.item(organization=self.other)
        result=json.loads(catalog.__wrapped__(self.request()).content)['data']
        self.assertEqual(result['total'],1);self.assertEqual(result['results'][0]['id'],str(item.id))
    @patch('apps.business.barcode_catalog.permitted',return_value=True)
    def test_export_preserves_numbers_as_text_and_neutralizes_formulas(self,_):
        self.item(barcode='000123456789',name='=1+1');self.item(organization=self.other,name='Hidden')
        response=export_catalog.__wrapped__(self.request())
        workbook=load_workbook(BytesIO(response.content));sheet=workbook.active
        self.assertEqual(sheet.max_row,2);self.assertEqual(sheet['C2'].value,'000123456789')
        self.assertEqual(sheet['B2'].value,'=1+1');self.assertEqual(sheet['B2'].data_type,'s')
    @patch('apps.business.barcode_catalog.permitted',return_value=True)
    def test_backfill_changes_only_empty_barcodes_in_current_tenant(self,_):
        old=self.item();old.update(set__barcode='')
        existing=self.item(barcode='00123');other=self.item(organization=self.other);other.update(set__barcode='')
        response=generate_missing.__wrapped__(self.request('post'))
        self.assertEqual(json.loads(response.content)['data']['updated'],1)
        old.reload();existing.reload();other.reload()
        self.assertEqual(old.barcode,Product.automatic_barcode(old.id));self.assertEqual(existing.barcode,'00123');self.assertEqual(other.barcode,'')
    @patch('apps.business.barcode_catalog.permitted',return_value=False)
    def test_permissions_block_export_and_backfill(self,_):
        self.assertEqual(export_catalog.__wrapped__(self.request()).status_code,403)
        self.assertEqual(generate_missing.__wrapped__(self.request('post')).status_code,403)

    def test_api_create_and_update_preserve_or_generate_barcode(self):
        from apps.products.services.product_api_service import ProductAPIService as S, ProductAPIValidationError
        payload=dict(sku='SCAN-TEST', name='Test product', category_id=str(self.category.id), unit='piece', barcode='000123456789')
        product=S.create_product(organization=self.org,payload=payload)
        self.assertEqual(product.barcode,'000123456789')
        with self.assertRaises(ProductAPIValidationError):
            S.create_product(organization=self.org,payload={**payload,'sku':'DUPLICATE'})
        other=self.item()
        with self.assertRaises(ProductAPIValidationError):
            S.update_product(organization=self.org,product_id=str(other.id),payload={'barcode':'000123456789'})
        changed=S.update_product(organization=self.org,product_id=str(product.id),payload={'barcode':''})
        self.assertEqual(changed.barcode,Product.automatic_barcode(product.id))
        automatic=S.create_product(organization=self.org,payload={**payload,'sku':'AUTO','barcode':''})
        self.assertEqual(automatic.barcode,Product.automatic_barcode(automatic.id))
