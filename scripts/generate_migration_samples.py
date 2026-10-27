"""Run after Django setup: python manage.py shell -c \"exec(open('scripts/generate_migration_samples.py').read())\""""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
from apps.migration.services.migration_schema_service import MigrationSchemaService
from apps.migration.services.template_examples import EXAMPLES
out=Path('migration_samples');out.mkdir(exist_ok=True)
for kind,schema in MigrationSchemaService.SCHEMAS.items():
    fields=schema['fields'];wb=Workbook();ws=wb.active;ws.title='Data';ws.append(list(fields));ws.freeze_panes='A2'
    ex=wb.create_sheet('Example only');ex.append(list(fields));ex.append([('Example worker' if k=='name' else EXAMPLES.get(k,'')) if kind=='EMPLOYEES' else EXAMPLES.get(k,'Example') for k in fields])
    guide=wb.create_sheet('Column guide');guide.append(['Column','Label','Required']);
    for k,v in fields.items():guide.append([k,v['label'],'Yes' if v.get('required') else 'No'])
    note=wb.create_sheet('Read me');note.append(['NexEase migration template']);note.append(['Fill the Data sheet. The Example only sheet is illustrative and is not imported.']);note.append(['Use existing codes/IDs from your own organisation. Import master records before transactions.']);note.append(['Payment imports affect actual recorded balances; validate and review first.']);note.append(['Sales invoices require fulfilled sales orders. Vendor bills require received purchase orders.']);note.append(['Use unique payment references; one allocation per row.'])
    for sheet in wb:
        sheet.freeze_panes='A2'
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='346B40')
        for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=min(60,max(18,max(len(str(c.value or '')) for c in col)+2))
    wb.save(out/(kind+'.xlsx'))
print('Created',len(MigrationSchemaService.SCHEMAS),'sample templates in',out)
